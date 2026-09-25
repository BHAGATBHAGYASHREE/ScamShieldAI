import os
import sys
import subprocess
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

TOTAL_SLIDES = 19
SLIDES_DIR = "exports/slides"
os.makedirs(SLIDES_DIR, exist_ok=True)

CHROME_PATH = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

print(f"[*] Starting capture of {TOTAL_SLIDES} slides...")

for i in range(1, TOTAL_SLIDES + 1):
    out_path = f"{SLIDES_DIR}/slide_{i}.png"
    url = f"http://localhost:8088/presentation.html?slide={i}&export=true"
    cmd = [
        CHROME_PATH,
        "--headless=new",
        f"--screenshot={out_path}",
        "--window-size=1920,1080",
        "--hide-scrollbars",
        url
    ]
    print(f"    Capturing Slide {i}/{TOTAL_SLIDES} -> {out_path}...", end=" ", flush=True)
    res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    if os.path.exists(out_path) and os.path.getsize(out_path) > 1000:
        print(f"Done ({os.path.getsize(out_path)} bytes)")
    else:
        print("FAILED!")
        print(res.stderr.decode('utf-8')[:300])

print("\n[*] Validating captured slides...")
images = []
for i in range(1, TOTAL_SLIDES + 1):
    path = f"{SLIDES_DIR}/slide_{i}.png"
    if not os.path.exists(path):
        raise FileNotFoundError(f"Missing slide: {path}")
    img = Image.open(path).convert("RGB")
    images.append(img)
print(f"    All {len(images)} slides verified successfully.")

# 1. Generate PDF
pdf_path = "ScamShield_AI_Presentation.pdf"
print(f"\n[*] Compiling PDF: {pdf_path}...")
images[0].save(pdf_path, save_all=True, append_images=images[1:], quality=95)
print(f"    PDF compiled successfully ({os.path.getsize(pdf_path) / (1024*1024):.2f} MB)")

# 2. Generate PPTX
pptx_path = "ScamShield_AI_Presentation.pptx"
print(f"\n[*] Compiling PPTX: {pptx_path}...")
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

for i, img_path in enumerate([f"{SLIDES_DIR}/slide_{j}.png" for j in range(1, TOTAL_SLIDES + 1)], start=1):
    slide = prs.slides.add_slide(blank_layout)
    slide.shapes.add_picture(
        img_path,
        left=0,
        top=0,
        width=prs.slide_width,
        height=prs.slide_height
    )

prs.save(pptx_path)
print(f"    PPTX compiled successfully ({os.path.getsize(pptx_path) / (1024*1024):.2f} MB)")

print("\n[SUCCESS] Both PDF and PPTX generated flawlessly!")
